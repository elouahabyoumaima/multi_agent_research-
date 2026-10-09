import io
import os
import re

import requests
import streamlit as st
from docx import Document
from fpdf import FPDF
from fpdf.enums import XPos, YPos
from fpdf.fonts import FontFace

API_URL = os.getenv("API_URL", "http://localhost:8000")
ACCENT = "#0ea5e9"  # bleu ciel. Pour du rouge: "#ef4444"

st.set_page_config(page_title="Research Assistant", page_icon="🔎", layout="centered")

# ---------------------------------------------------------------------------
# Etat de session
# ---------------------------------------------------------------------------
if "token" not in st.session_state:
    st.session_state.token = None
if "username" not in st.session_state:
    st.session_state.username = None
if "theme" not in st.session_state:
    st.session_state.theme = "light"


def auth_headers():
    return {"Authorization": "Bearer " + st.session_state.token}


# ---------------------------------------------------------------------------
# Theme clair / sombre
# ---------------------------------------------------------------------------
PALETTES = {
    "light": {
        "bg": "#FFFFFF", "text": "#111111", "muted": "#6B7280",
        "sidebar": "#F7F7F8", "border": "#E5E7EB", "input": "#FFFFFF",
        "btn": "#FFFFFF", "th": "#F3F4F6",
    },
    "dark": {
        "bg": "#212121", "text": "#ECECEC", "muted": "#9CA3AF",
        "sidebar": "#171717", "border": "#3A3A3A", "input": "#2F2F2F",
        "btn": "#2F2F2F", "th": "#2F2F2F",
    },
}

CSS_TEMPLATE = """
<style>
    .stApp, [data-testid="stHeader"] { background-color: __BG__ !important; }
    #MainMenu, footer { visibility: hidden; }

    /* tout le texte suit le theme: noir sur fond blanc, blanc sur fond noir */
    .stApp, .stApp p, .stApp span, .stApp li, .stApp td, .stApp th, .stApp label,
    .stApp h1, .stApp h2, .stApp h3, .stApp h4, .stApp h5, .stApp h6,
    [data-testid="stMarkdownContainer"], [data-testid="stMarkdownContainer"] *,
    [data-testid="stWidgetLabel"] p, [data-testid="stSidebar"] *,
    button[data-baseweb="tab"], button[data-baseweb="tab"] * {
        color: __TEXT__;
    }
    .stApp .step-badge { color: __ACCENT__ !important; }
    [data-testid="stHeader"] *, [data-testid="stToolbar"] * { color: __TEXT__ !important; }
    [data-testid="stCaptionContainer"], [data-testid="stCaptionContainer"] p { color: __MUTED__ !important; }

    .main-header { text-align: center; padding: 1.2rem 0 1.5rem 0; }
    .main-header h1 { color: __TEXT__; font-size: 1.6rem; font-weight: 600; margin-bottom: 0.2rem; }
    .main-header p { color: __MUTED__; font-size: 0.9rem; }

    .stApp .step-badge {
        display: inline-block;
        background-color: __ACCENT__1A;
        color: __ACCENT__ !important;
        padding: 3px 12px;
        border-radius: 20px;
        font-size: 0.8rem;
        font-weight: 600;
        margin: 2px 4px 2px 0;
    }

    section[data-testid="stSidebar"] {
        background-color: __SIDEBAR__ !important;
        border-right: 1px solid __BORDER__;
    }

    /* champs */
    div[data-baseweb="input"] > div, div[data-baseweb="textarea"], div[data-baseweb="base-input"] {
        background-color: __INPUT__ !important;
        border-color: __BORDER__ !important;
    }
    input, textarea { color: __TEXT__ !important; -webkit-text-fill-color: __TEXT__ !important; background-color: __INPUT__ !important; }
    input::placeholder, textarea::placeholder { color: __MUTED__ !important; -webkit-text-fill-color: __MUTED__ !important; }
    [data-testid="stForm"] { border-color: __BORDER__ !important; }

    /* onglets */
    button[data-baseweb="tab"] p { color: __TEXT__; }

    /* boutons secondaires + telechargement */
    .stButton > button:not([kind="primary"]), .stDownloadButton > button {
        background-color: __BTN__ !important;
        border: 1px solid __BORDER__ !important;
    }
    .stButton > button:not([kind="primary"]) p, .stDownloadButton > button p { color: __TEXT__ !important; }

    /* bouton principal */
    .stButton > button[kind="primary"], [data-testid="stFormSubmitButton"] > button[kind="primary"],
    button[kind="primaryFormSubmit"] {
        background-color: __ACCENT__ !important;
        border: 1px solid __ACCENT__ !important;
    }
    .stButton > button[kind="primary"] p, button[kind="primaryFormSubmit"] p { color: #FFFFFF !important; }

    /* tableaux markdown */
    .stMarkdown table { border-collapse: collapse; width: 100%; }
    .stMarkdown th { background-color: __TH__; }
    .stMarkdown th, .stMarkdown td { border: 1px solid __BORDER__; padding: 6px 10px; }
    .stMarkdown a { color: __ACCENT__; }
    hr { border-color: __BORDER__ !important; }
</style>
"""


def apply_theme():
    p = PALETTES[st.session_state.theme]
    css = CSS_TEMPLATE
    for key, value in {
        "__BG__": p["bg"], "__TEXT__": p["text"], "__MUTED__": p["muted"],
        "__SIDEBAR__": p["sidebar"], "__BORDER__": p["border"], "__INPUT__": p["input"],
        "__BTN__": p["btn"], "__TH__": p["th"], "__ACCENT__": ACCENT,
    }.items():
        css = css.replace(key, value)
    st.markdown(css, unsafe_allow_html=True)


def theme_button(key):
    label = "🌙 Mode sombre" if st.session_state.theme == "light" else "☀️ Mode clair"
    if st.button(label, key=key, use_container_width=True):
        st.session_state.theme = "dark" if st.session_state.theme == "light" else "light"
        st.rerun()


apply_theme()


# ---------------------------------------------------------------------------
# Export PDF (tableaux, titres, listes, gras)
# ---------------------------------------------------------------------------
FONT_SETS = [
    ("C:/Windows/Fonts/arial.ttf", "C:/Windows/Fonts/arialbd.ttf", "C:/Windows/Fonts/ariali.ttf"),
    ("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
     "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf",
     "/usr/share/fonts/truetype/dejavu/DejaVuSans-Oblique.ttf"),
]

CHAR_FIXES = {
    "\u2011": "-", "\u2010": "-", "\u2013": "-", "\u2014": "-",
    "\u202f": " ", "\u00a0": " ", "\u2019": "'", "\u2018": "'",
    "\u201c": '"', "\u201d": '"', "\u2026": "...", "\u2192": "->",
    "\u2022": "-",
}


def _setup_font(pdf):
    """Charge une police Unicode (Arial sur Windows). Sinon Helvetica."""
    for regular, bold, italic in FONT_SETS:
        if os.path.exists(regular) and os.path.exists(bold):
            pdf.add_font("Main", "", regular)
            pdf.add_font("Main", "B", bold)
            pdf.add_font("Main", "I", italic if os.path.exists(italic) else regular)
            pdf.add_font("Main", "BI", bold)
            return "Main", True
    return "Helvetica", False


def _clean_inline(s):
    s = re.sub(r"\[([^\]]+)\]\((https?://[^)]+)\)", r"\1 (\2)", s)  # liens
    s = s.replace("`", "")
    s = re.sub(r"(?<!\*)\*(?!\*)", "", s)  # etoile seule supprimee, ** reste pour le gras
    return s.strip()


def _prepare(text, unicode_ok):
    for old, new in CHAR_FIXES.items():
        text = text.replace(old, new)
    text = re.sub(r"<br\s*/?>", "\n", text, flags=re.IGNORECASE)
    if not unicode_ok:
        text = text.encode("latin-1", "replace").decode("latin-1")
    return text


def _render_table(pdf, font, rows):
    ncols = max(len(r) for r in rows)
    rows = [r + [""] * (ncols - len(r)) for r in rows]

    weights = []
    for c in range(ncols):
        longest = max(len(r[c]) for r in rows)
        weights.append(min(max(longest, 8), 45))

    pdf.set_font(font, "", 9)
    pdf.set_x(pdf.l_margin)
    with pdf.table(
        col_widths=tuple(weights),
        text_align="LEFT",
        line_height=5,
        padding=1.5,
        first_row_as_headings=True,
        headings_style=FontFace(emphasis="BOLD", color=(255, 255, 255), fill_color=(14, 165, 233)),
        markdown=True,
    ) as table:
        for r in rows:
            row = table.row()
            for cell in r:
                row.cell(_clean_inline(cell))
    pdf.ln(4)
    pdf.set_x(pdf.l_margin)


def report_to_pdf(report_text):
    pdf = FPDF()
    pdf.set_margins(15, 15, 15)
    pdf.set_auto_page_break(auto=True, margin=15)
    pdf.add_page()
    font, unicode_ok = _setup_font(pdf)

    text = _prepare(report_text, unicode_ok)
    lines = text.split("\n")
    usable = pdf.w - pdf.l_margin - pdf.r_margin
    heading_sizes = {1: 17, 2: 14, 3: 12, 4: 11}

    i = 0
    while i < len(lines):
        stripped = lines[i].strip()

        # ---- tableau markdown ----
        if stripped.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                if not all(re.fullmatch(r":?-{2,}:?", c) for c in cells if c):
                    rows.append(cells)
                i += 1
            if rows:
                _render_table(pdf, font, rows)
            continue

        # ---- ligne vide ----
        if not stripped:
            pdf.ln(3)
            i += 1
            continue

        # ---- separateur horizontal ----
        if re.fullmatch(r"-{3,}|\*{3,}|_{3,}", stripped):
            y = pdf.get_y() + 2
            pdf.set_draw_color(200, 200, 200)
            pdf.line(pdf.l_margin, y, pdf.w - pdf.r_margin, y)
            pdf.ln(5)
            i += 1
            continue

        # ---- titres ----
        m = re.match(r"^(#{1,6})\s+(.*)", stripped)
        if m:
            level = min(len(m.group(1)), 4)
            title = _clean_inline(m.group(2)).replace("**", "")
            pdf.ln(3 if level > 1 else 2)
            pdf.set_font(font, "B", heading_sizes[level])
            if level <= 2:
                pdf.set_text_color(14, 165, 233)
            else:
                pdf.set_text_color(17, 17, 17)
            pdf.set_x(pdf.l_margin)
            pdf.multi_cell(usable, heading_sizes[level] * 0.55, title,
                           new_x=XPos.LMARGIN, new_y=YPos.NEXT)
            if level == 1:
                y = pdf.get_y() + 1
                pdf.set_draw_color(14, 165, 233)
                pdf.line(pdf.l_margin, y, pdf.w - pdf.r_margin, y)
                pdf.ln(4)
            else:
                pdf.ln(1.5)
            pdf.set_text_color(0, 0, 0)
            i += 1
            continue

        # ---- listes a puces ----
        m = re.match(r"^[-*+]\s+(.*)", stripped)
        if m:
            pdf.set_font(font, "", 10)
            pdf.set_x(pdf.l_margin + 4)
            pdf.multi_cell(usable - 4, 5.5, "-  " + _clean_inline(m.group(1)),
                           markdown=True, new_x=XPos.LMARGIN, new_y=YPos.NEXT)
            i += 1
            continue

        # ---- listes numerotees ----
        m = re.match(r"^(\d+[.)])\s+(.*)", stripped)
        if m:
            pdf.set_font(font, "", 10)
            pdf.set_x(pdf.l_margin + 4)
            pdf.multi_cell(usable - 4, 5.5, m.group(1) + "  " + _clean_inline(m.group(2)),
                           markdown=True, new_x=XPos.LMARGIN, new_y=YPos.NEXT)
            i += 1
            continue

        # ---- citation ">" ----
        if stripped.startswith(">"):
            stripped = stripped.lstrip("> ").strip()
            pdf.set_text_color(90, 90, 90)

        # ---- paragraphe normal ----
        pdf.set_font(font, "", 10)
        pdf.set_x(pdf.l_margin)
        pdf.multi_cell(usable, 5.5, _clean_inline(stripped),
                       markdown=True, new_x=XPos.LMARGIN, new_y=YPos.NEXT)
        pdf.set_text_color(0, 0, 0)
        i += 1

    return bytes(pdf.output())


# ---------------------------------------------------------------------------
# Export Word / TXT
# ---------------------------------------------------------------------------
def report_to_docx(report_text):
    doc = Document()
    for line in report_text.split("\n"):
        doc.add_paragraph(line)
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def render_export_buttons(report_text, key_prefix):
    col1, col2, col3 = st.columns(3)
    with col1:
        st.download_button(
            "📄 TXT",
            report_text,
            file_name="rapport.txt",
            key=key_prefix + "_txt",
            use_container_width=True,
        )
    with col2:
        st.download_button(
            "📕 PDF",
            report_to_pdf(report_text),
            file_name="rapport.pdf",
            mime="application/pdf",
            key=key_prefix + "_pdf",
            use_container_width=True,
        )
    with col3:
        docx_mime = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        st.download_button(
            "📘 Word",
            report_to_docx(report_text),
            file_name="rapport.docx",
            mime=docx_mime,
            key=key_prefix + "_docx",
            use_container_width=True,
        )


# ---------------------------------------------------------------------------
# Page de connexion / inscription
# ---------------------------------------------------------------------------
def login_page():
    _, col_theme = st.columns([3, 1])
    with col_theme:
        theme_button("theme_login")

    header_html = (
        '<div class="main-header">'
        '<h1>Multi-Agent Research Assistant</h1>'
        '<p>Connecte-toi pour acceder a l assistant et ton historique</p>'
        '</div>'
    )
    st.markdown(header_html, unsafe_allow_html=True)

    tab_login, tab_signup = st.tabs(["Connexion", "Creer un compte"])

    # ---------------- Connexion ----------------
    with tab_login:
        with st.form("login_form"):
            username = st.text_input("Nom d'utilisateur", key="login_username")
            password = st.text_input("Mot de passe", type="password", key="login_password")
            submitted = st.form_submit_button("Se connecter", use_container_width=True, type="primary")
            if submitted:
                try:
                    resp = requests.post(
                        API_URL + "/auth/login",
                        json={"username": username, "password": password},
                    )
                    if resp.status_code == 200:
                        st.session_state.token = resp.json()["access_token"]
                        st.session_state.username = username
                        st.rerun()
                    else:
                        st.error(resp.json().get("detail", "Erreur de connexion"))
                except Exception as e:
                    st.error("Erreur : " + str(e))

    # ---------------- Inscription + verification ----------------
    with tab_signup:
        pending_email = st.session_state.get("pending_email")

        if pending_email:
            st.info("Un code a 6 chiffres a ete envoye a : " + pending_email)
            with st.form("verify_form"):
                code = st.text_input("Code de verification", max_chars=6)
                verify_clicked = st.form_submit_button("Verifier", use_container_width=True, type="primary")
                if verify_clicked:
                    try:
                        resp = requests.post(
                            API_URL + "/auth/verify",
                            json={"email": pending_email, "code": code},
                        )
                        if resp.status_code == 200:
                            st.session_state.token = resp.json()["access_token"]
                            st.session_state.username = st.session_state.get("pending_username")
                            st.session_state.pending_email = None
                            st.session_state.pending_username = None
                            st.rerun()
                        else:
                            st.error(resp.json().get("detail", "Code incorrect"))
                    except Exception as e:
                        st.error("Erreur : " + str(e))

            col_a, col_b = st.columns(2)
            with col_a:
                if st.button("Renvoyer le code", use_container_width=True):
                    try:
                        resp = requests.post(API_URL + "/auth/resend", json={"email": pending_email})
                        if resp.status_code == 200 and resp.json().get("email_sent"):
                            st.success("Nouveau code envoye.")
                        else:
                            st.error("Echec de l'envoi de l'email.")
                    except Exception as e:
                        st.error("Erreur : " + str(e))
            with col_b:
                if st.button("Annuler", use_container_width=True):
                    st.session_state.pending_email = None
                    st.session_state.pending_username = None
                    st.rerun()
        else:
            with st.form("signup_form"):
                username = st.text_input("Nom d'utilisateur", key="signup_username")
                email = st.text_input("Email", key="signup_email")
                phone = st.text_input("Telephone (optionnel)", key="signup_phone")
                password = st.text_input("Mot de passe", type="password", key="signup_password")
                password_confirm = st.text_input("Confirmer le mot de passe", type="password", key="signup_password_confirm")
                submitted = st.form_submit_button("Creer le compte", use_container_width=True, type="primary")
                if submitted:
                    try:
                        payload = {
                            "username": username,
                            "email": email,
                            "phone": phone or None,
                            "password": password,
                            "password_confirm": password_confirm,
                        }
                        resp = requests.post(API_URL + "/auth/signup", json=payload)
                        if resp.status_code == 200:
                            if not resp.json().get("email_sent"):
                                st.warning("Compte cree, mais l'email n'a pas pu etre envoye. Utilise 'Renvoyer le code'.")
                            st.session_state.pending_email = email
                            st.session_state.pending_username = username
                            st.rerun()
                        else:
                            detail = resp.json().get("detail", "Erreur lors de la creation")
                            st.error(detail if isinstance(detail, str) else "Donnees invalides (verifie l'email et les champs).")
                    except Exception as e:
                        st.error("Erreur : " + str(e))


# ---------------------------------------------------------------------------
# Page principale (apres connexion)
# ---------------------------------------------------------------------------
def main_page():
    header_html = (
        '<div class="main-header">'
        '<h1>Multi-Agent Research Assistant</h1>'
        '<p>Planification -> Recherche (web + MCP) -> Synthese</p>'
        '</div>'
    )
    st.markdown(header_html, unsafe_allow_html=True)

    with st.sidebar:
        st.markdown("### Utilisateur : " + st.session_state.username)
        theme_button("theme_main")
        if st.button("Se deconnecter", use_container_width=True):
            st.session_state.token = None
            st.session_state.username = None
            st.rerun()

        st.markdown("---")
        st.markdown("### Historique")
        try:
            resp = requests.get(API_URL + "/history", headers=auth_headers())
            history = resp.json() if resp.status_code == 200 else []
        except Exception:
            history = []

        if not history:
            st.caption("Aucune recherche pour l'instant.")
        else:
            for item in history:
                label = item["question"][:40]
                if len(item["question"]) > 40:
                    label = label + "..."
                if st.button(label, key="hist_" + str(item["id"]), use_container_width=True):
                    st.session_state.viewing = item
                    st.rerun()

    # affichage d'un element d'historique selectionne
    if st.session_state.get("viewing"):
        item = st.session_state.viewing
        caption_text = "Recherche du " + item["created_at"][:10] + " -- " + str(item["duration_seconds"]) + "s"
        st.caption(caption_text)
        if st.button("Retour / Nouvelle recherche"):
            st.session_state.viewing = None
            st.rerun()
        st.markdown("---")
        st.markdown(item["report"])
        render_export_buttons(item["report"], "hist_" + str(item["id"]))
        return

    question = st.text_area(
        "Pose ta question de recherche :",
        height=80,
        placeholder="Ex: Quels sont les avantages de Rust par rapport a Go ?",
    )

    if st.button("Lancer la recherche", use_container_width=True, type="primary"):
        if not question.strip():
            st.warning("Ecris une question d'abord.")
        else:
            with st.spinner("Planification, recherche et redaction en cours (1-2 minutes)..."):
                try:
                    url = API_URL + "/research"
                    payload = {"question": question}
                    headers = auth_headers()
                    resp = requests.post(url, json=payload, headers=headers, timeout=180)
                    resp.raise_for_status()
                    result = resp.json()
                except Exception as e:
                    st.error("Erreur : " + str(e))
                    result = None

            if result:
                st.session_state.last_result = result

    if st.session_state.get("last_result") and not st.session_state.get("viewing"):
        result = st.session_state.last_result
        st.caption("Termine en " + str(result["duration_seconds"]) + "s")
        for sq in result["sub_questions"]:
            badge_html = '<span class="step-badge">' + sq[:60] + "...</span>"
            st.markdown(badge_html, unsafe_allow_html=True)
        st.markdown("---")
        st.markdown(result["report"])
        render_export_buttons(result["report"], "last")


# ---------------------------------------------------------------------------
# Routage
# ---------------------------------------------------------------------------
if st.session_state.token is None:
    login_page()
else:
    main_page()